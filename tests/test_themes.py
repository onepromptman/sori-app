"""themes/: the studio default, and Sori surviving without the private engine.

`engine/` holds a vendored private renderer and licensed fonts. A public
checkout does not ship it, so every one of these tests is really asking the
same question: does Sori still produce a deck when the branded backend is
simply not there?
"""

from __future__ import annotations

import deck_render
import pathlib

import pytest
from themes import studio

ALL_KINDS = [
    {"kind": "cover", "title": "T", "lede": "l", "eyebrow": "E"},
    {"kind": "divider", "title": "D", "eyebrow": "E", "lede": "l"},
    {"kind": "content", "heading": "H", "points": [{"label": "a", "sub": "b"}]},
    {
        "kind": "content_stat",
        "heading": "H",
        "hero_stat": {"num": "90", "label": "m", "desc": "d"},
        "points": [{"label": "a", "sub": "b"}],
    },
    {
        "kind": "two_column",
        "heading": "H",
        "lede": "l",
        "left": {"heading": "L", "points": [{"label": "a"}]},
        "right": {"heading": "R", "points": [{"label": "b"}]},
    },
    {"kind": "process", "heading": "H", "steps": [{"label": "a", "sub": "b"}], "dominant": 0},
    {"kind": "agenda", "title": "H", "items": [{"label": "a", "sub": "b"}], "active": 0},
    {"kind": "quote", "quote": "q", "attribution": "who"},
    {"kind": "callout", "heading": "H", "lede": "l"},
    {"kind": "full_bleed", "heading": "H", "lede": "l", "eyebrow": "E"},
    {
        "kind": "data",
        "hero_stat": {"num": "90", "label": "m", "desc": "d"},
        "rail_label": "R",
        "stats": [{"num": "3", "label": "x", "desc": "y"}],
    },
    {"kind": "scorecard", "heading": "H", "rows": [{"label": "r", "level": 0, "note": "n"}]},
    {
        "kind": "chart",
        "heading": "H",
        "unit": "%",
        "highlight": 0,
        "data": [{"label": "a", "value": 3.0}, {"label": "b", "value": 2.0}, {"label": "c", "value": 1.0}],
    },
    {"kind": "closing", "heading": "H", "cta": "c", "eyebrow": "E"},
]


def test_studio_renders_every_kind(tmp_path):
    """A theme that cannot paint one kind would crash mid-deck at a user's desk."""
    assert {s["kind"] for s in ALL_KINDS} == set(studio.WRITERS), "ALL_KINDS drifted from WRITERS"
    out = tmp_path / "d.pptx"
    studio.render(ALL_KINDS, str(out))
    assert out.exists() and out.stat().st_size > 0


def test_studio_rejects_an_unknown_kind(tmp_path):
    with pytest.raises(KeyError):
        studio.render([{"kind": "hologram"}], str(tmp_path / "d.pptx"))


def test_notes_survive_into_the_deck(tmp_path):
    from pptx import Presentation

    out = tmp_path / "d.pptx"
    studio.render([{"kind": "content", "heading": "H", "notes": "speaker note"}], str(out))
    prs = Presentation(str(out))
    assert prs.slides[0].notes_slide.notes_text_frame.text == "speaker note"


def test_studio_is_always_offered_and_the_engine_is_optional():
    assert "studio" in deck_render.THEMES
    assert deck_render.DEFAULT_THEME in deck_render.THEMES


def test_without_the_private_engine_sori_still_renders(tmp_path, monkeypatch):
    """The public-checkout case: engine/ absent, studio carries the deck."""
    monkeypatch.setattr(deck_render, "deckgen", None)
    monkeypatch.setattr(deck_render, "THEMES", {"studio": "studio"})
    outline = {
        "title": "T",
        "subtitle": "s",
        "closing_heading": "C",
        "cta": "go",
        "slides": [{"kind": "content", "heading": "H", "points": [{"label": "a", "sub": "b"}]}],
    }
    out = deck_render.render_pptx(outline, "studio", tmp_path / "d.pptx")
    assert out.exists()
    # and a branded theme is refused rather than silently falling back
    with pytest.raises(deck_render.DeckRenderError):
        deck_render.render_pptx(outline, "baek", tmp_path / "e.pptx")


def test_two_column_carries_its_heading():
    """Regression: heading/lede were dropped, so the slide rendered as two
    unlabelled panes with no statement of what was being compared."""
    spec = deck_render._slide_spec(
        {
            "kind": "two_column",
            "heading": "H",
            "lede": "l",
            "eyebrow": "E",
            "left": {"heading": "L"},
            "right": {"heading": "R"},
        },
        "studio",
    )
    assert spec["heading"] == "H" and spec["lede"] == "l"


@pytest.mark.parametrize(
    "text, width, pt, expected_min",
    [
        ("short", 6.0, 13, 1),
        ("x" * 400, 6.0, 13, 4),
    ],
)
def test_wrapped_lines_grows_with_text(text, width, pt, expected_min):
    """Row height must track wrapping, or a long label overlaps its own sub."""
    assert studio._wrapped_lines(text, width, pt) >= expected_min


def test_studio_names_only_stock_faces(tmp_path):
    """OOXML has no fallback chain and python-pptx cannot embed a face, so a
    name the opening machine lacks becomes PowerPoint's default, not our
    second choice. studio may therefore only name fonts that are already on a
    stock install -- no font installation, on any platform."""
    assert studio.FONT_DISPLAY == "Georgia", "Georgia ships on Windows and macOS"
    for name in (studio.FONT_BODY, studio.FONT_MONO, studio.FONT_NUMERAL):
        assert name in {"Segoe UI", "Consolas"}, f"{name} is not a stock face"
    out = tmp_path / "d.pptx"
    studio.render([{"kind": "content", "heading": "H", "points": [{"label": "a", "sub": "b"}]}], str(out))
    assert out.exists()


def test_a_public_render_installs_no_fonts(tmp_path, monkeypatch):
    """Regression: the studio path must never reach the font installer."""
    called = []
    if deck_render.install_fonts is not None:
        monkeypatch.setattr(deck_render.install_fonts, "ensure_installed", lambda: called.append(1) or True)
    monkeypatch.setattr(deck_render, "deckgen", None)
    monkeypatch.setattr(deck_render, "THEMES", {"studio": "studio"})
    outline = {
        "title": "T",
        "subtitle": "s",
        "closing_heading": "C",
        "cta": "go",
        "slides": [{"kind": "content", "heading": "H"}],
    }
    deck_render.render_pptx(outline, "studio", tmp_path / "d.pptx")
    assert called == [], "studio must not install fonts"


def test_every_bundled_font_is_open_licensed():
    """The bundled faces serve the optional private engine, which names them
    directly. They ship only while every one is OFL and the licence ships
    with them. Absent in a public checkout -- skip rather than fail."""
    if deck_render.install_fonts is None:
        pytest.skip("no engine/ in this checkout, so no bundled fonts")
    fonts = deck_render.install_fonts.FONTS_DIR
    assert (fonts / "OFL.txt").is_file(), "OFL 1.1 text must ship with the fonts"
    licence = (fonts / "OFL.txt").read_text(encoding="utf-8")
    assert "SIL OPEN FONT LICENSE Version 1.1" in licence
    for holder in ("Gowun Batang Project Authors", "IBM Corp"):
        assert holder in licence, f"{holder} copyright notice missing"
    assert list(fonts.glob("*.ttf")), "no fonts bundled"


def test_advance_constants_are_not_optimistic():
    """The layout estimates are only safe while the constants stay at or above
    what the named faces actually measure. Pillow is not a declared
    dependency, so this runs where it exists and skips where it does not."""
    ImageFont = pytest.importorskip("PIL.ImageFont")
    samples = [
        "GPT-6 Astra: the claim and the receipts",
        "Capability demos are not adoption",
        "When Three Frontier Models Go Dark at Once",
    ]
    cases = [
        ("C:/Windows/Fonts/georgiab.ttf", studio.ADV_DISPLAY),
        ("C:/Windows/Fonts/segoeui.ttf", studio.ADV_BODY),
    ]
    for path, constant in cases:
        if not pathlib.Path(path).is_file():
            pytest.skip(f"{path} not present on this platform")
        f = ImageFont.truetype(path, 200)
        worst = max(f.getlength(s) / 200 / len(s) for s in samples)
        assert constant >= worst, (
            f"{path} measures {worst:.3f} em/char but the layout assumes "
            f"{constant}; under-estimating overprints the next block"
        )


@pytest.mark.parametrize(
    "text, expected",
    [
        ("Concentration\nrisk", 2),
        ("When Three Frontier Models\nGo Dark at Once", 2),
        ("a\nb\nc", 3),
        ("no breaks here", 1),
    ],
)
def test_hard_line_breaks_are_counted(text, expected):
    """An outline controls where a headline turns by writing a hard break.
    Measuring that as one long line overprints whatever is placed under it."""
    assert studio._wrapped_lines(text, 12.0, 20.0) == expected
